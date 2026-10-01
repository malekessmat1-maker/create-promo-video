"""Builds Sorted-Budget-Planner.xlsx: the product a customer downloads.
Every number a customer sees is a formula over their own inputs (yellow cells).
Run: python3 sorted/scripts/build_planner.py sorted/product/Sorted-Budget-Planner.xlsx
"""
import sys, random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "Sorted-Budget-Planner.xlsx"
YEAR = 2027
INK, BRAND, MINT, INPUT, LINE, MUTED, CORAL = "1E2A2F", "2E7D5B", "E8F3EE", "FFF6D6", "D5DED9", "6B7C76", "E0603F"
F = lambda **k: Font(name="Arial", **{"size": 10, "color": INK, **k})
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=LINE)
BOX = Border(bottom=thin)
MONEY = '#,##0.00;-#,##0.00;"-"'
PCT = '0%;-0%;"-"'
MONTHS = ["January","February","March","April","May","June","July","August","September","October","November","December"]
GROUPS = ["Income", "Needs", "Wants", "Savings", "Debt"]

wb = Workbook()

def sheet(title, tab, widths):
    ws = wb.create_sheet(title) if wb.active.title != "Sheet" else wb.active
    if ws.title == "Sheet": ws.title = title
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    return ws

def title(ws, text, sub):
    ws["B2"] = text; ws["B2"].font = F(size=20, bold=True, color=BRAND)
    ws["B3"] = sub; ws["B3"].font = F(size=10, color=MUTED)
    ws.row_dimensions[2].height = 30

def header(ws, row, col0, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = F(bold=True, color="FFFFFF"); c.fill = fill(BRAND)
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center")
    ws.row_dimensions[row].height = 20

def inp(c, fmt=None):
    c.fill = fill(INPUT); c.font = F(color="1F3A93"); c.border = BOX
    if fmt: c.number_format = fmt

def out(c, fmt=None, bold=False):
    c.font = F(bold=bold); c.border = BOX
    if fmt: c.number_format = fmt

# ------------------------------------------------------------------ Start Here
ws = sheet("Start Here", BRAND, {"A": 3, "B": 4, "C": 90})
ws["C2"] = "Sorted. Budget Planner"; ws["C2"].font = F(size=24, bold=True, color=BRAND)
ws["C3"] = "Ten minutes a week. Every pound, dollar or euro with a job."; ws["C3"].font = F(size=11, color=MUTED)
steps = [
    ("1", "Setup tab", "Type your budget year, your income sources and what you plan to spend per category each month. Rename or add categories freely (up to 30)."),
    ("2", "Transactions tab", "Log each payment and paycheck: date, description, pick a category from the dropdown, amount. The example rows are there to show the format: delete them when you start."),
    ("3", "This Month tab", "Pick a month from the dropdown. See budget vs actual for every category, what's left, and your Needs / Wants / Savings split against 50/30/20."),
    ("4", "Year Dashboard tab", "Twelve months at a glance with charts: income, spending by group, savings rate."),
    ("5", "Debt Payoff tab", "List your debts in the order you want to clear them, add any extra monthly payment, and see your debt-free date and total interest."),
    ("6", "Savings Goals tab", "Set targets and dates. Sorted tells you how much to put aside each month."),
]
r = 6
ws.cell(row=5, column=3, value="How to use it").font = F(size=13, bold=True)
for n, head, body in steps:
    ws.cell(row=r, column=2, value=n).font = F(size=12, bold=True, color="FFFFFF")
    ws.cell(row=r, column=2).fill = fill(BRAND); ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="top")
    c = ws.cell(row=r, column=3, value=f"{head}: {body}")
    c.font = F(size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32; r += 1
r += 1
ws.cell(row=r, column=3, value="Colour key").font = F(size=13, bold=True); r += 1
ws.cell(row=r, column=2).fill = fill(INPUT)
ws.cell(row=r, column=3, value="Yellow cells are yours to type in. Everything else calculates itself, so you can't break it.").font = F(); r += 2
ws.cell(row=r, column=3, value="Works in").font = F(size=13, bold=True); r += 1
ws.cell(row=r, column=3, value="Microsoft Excel (2010 or newer, Windows and Mac), Google Sheets (File > Import > Upload) and Apple Numbers.").font = F(); r += 1
ws.cell(row=r, column=3, value="Amounts have no currency symbol, so the planner works in any currency.").font = F(); r += 2
ws.cell(row=r, column=3, value="Tip: back up a blank copy before you start, so you can reuse it next year.").font = F(italic=True, color=MUTED)

# ------------------------------------------------------------------ Setup
ws = sheet("Setup", "5BA67F", {"A": 3, "B": 30, "C": 14, "D": 18, "E": 3, "F": 34, "G": 16})
title(ws, "Setup", "Set this once. Change it whenever life changes.")
ws["B5"] = "Budget year"; ws["B5"].font = F(bold=True)
ws["C5"] = YEAR; inp(ws["C5"], "0")
header(ws, 7, 2, ["Category", "Group", "Monthly budget"])
cats = [("Salary", "Income", 4200), ("Side income", "Income", 400),
        ("Rent / mortgage", "Needs", 1350), ("Utilities", "Needs", 180), ("Groceries", "Needs", 450), ("Transport", "Needs", 220),
        ("Insurance", "Needs", 160), ("Phone & internet", "Needs", 90), ("Health", "Needs", 60),
        ("Eating out", "Wants", 200), ("Shopping", "Wants", 150), ("Entertainment", "Wants", 80), ("Subscriptions", "Wants", 45),
        ("Travel", "Wants", 120), ("Gifts", "Wants", 50),
        ("Emergency fund", "Savings", 500), ("Investing", "Savings", 445), ("Holiday fund", "Savings", 100),
        ("Credit card", "Debt", 250), ("Student loan", "Debt", 150)]
CAT0, CAT1 = 8, 37  # 30 category slots
for i in range(CAT0, CAT1 + 1):
    for col in (2, 3, 4): inp(ws.cell(row=i, column=col), MONEY if col == 4 else None)
for i, (n, g, b) in enumerate(cats):
    ws.cell(row=CAT0 + i, column=2, value=n); ws.cell(row=CAT0 + i, column=3, value=g); ws.cell(row=CAT0 + i, column=4, value=b)
dv = DataValidation(type="list", formula1='"' + ",".join(GROUPS) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"C{CAT0}:C{CAT1}")
CATS = f"Setup!$B${CAT0}:$B${CAT1}"; GRPS = f"Setup!$C${CAT0}:$C${CAT1}"; BUDS = f"Setup!$D${CAT0}:$D${CAT1}"
ws["F7"] = "Your plan in one look"; ws["F7"].font = F(size=12, bold=True)
plan = [("Planned income", f'=SUMIFS($D${CAT0}:$D${CAT1},$C${CAT0}:$C${CAT1},"Income")')]
for g in GROUPS[1:]:
    plan.append((f"Planned {g.lower()}", f'=SUMIFS($D${CAT0}:$D${CAT1},$C${CAT0}:$C${CAT1},"{g}")'))
plan.append(("Left to assign", "=G8-SUM(G9:G12)"))
for i, (lab, f) in enumerate(plan):
    ws.cell(row=8 + i, column=6, value=lab).font = F(bold=(lab == "Left to assign"))
    out(ws.cell(row=8 + i, column=7, value=f), MONEY, bold=(lab == "Left to assign"))
ws["F14"] = "Aim for 0: every unit of income gets a job (zero-based budgeting)."
ws["F14"].font = F(size=9, italic=True, color=MUTED); ws["F14"].alignment = Alignment(wrap_text=True); ws.row_dimensions[14].height = 26
ws.conditional_formatting.add("G13", CellIsRule(operator="lessThan", formula=["-0.005"], font=Font(name="Arial", bold=True, color=CORAL)))
ws.conditional_formatting.add("G13", CellIsRule(operator="between", formula=["-0.005", "0.005"], font=Font(name="Arial", bold=True, color=BRAND)))
ws.freeze_panes = "B8"

# ------------------------------------------------------------------ Transactions
ws = sheet("Transactions", "E0A33F", {"A": 3, "B": 13, "C": 32, "D": 22, "E": 14, "F": 12, "G": 8, "H": 8})
title(ws, "Transactions", "Log money in and money out. Pick categories from the dropdown. The grey columns fill themselves.")
header(ws, 4, 2, ["Date", "Description", "Category", "Amount", "Group", "Month", "Year"])
T0, T1 = 5, 1504  # 1,500 rows
rnd = random.Random(42)
ex = []
for m in (1, 2, 3):
    d = lambda day: dt.date(YEAR, m, day)
    ex += [(d(1), "Monthly salary", "Salary", 4200), (d(1), "Rent", "Rent / mortgage", 1350), (d(2), "Emergency fund transfer", "Emergency fund", 500),
           (d(2), "Index fund", "Investing", 445), (d(3), "Electric & water", "Utilities", round(rnd.uniform(150, 195), 2)),
           (d(4), "Weekly shop", "Groceries", round(rnd.uniform(95, 130), 2)), (d(5), "Train pass", "Transport", 180),
           (d(6), "Streaming + music", "Subscriptions", 45), (d(8), "Dinner with friends", "Eating out", round(rnd.uniform(40, 85), 2)),
           (d(11), "Weekly shop", "Groceries", round(rnd.uniform(95, 130), 2)), (d(12), "Phone + broadband", "Phone & internet", 90),
           (d(14), "Freelance logo job", "Side income", [350, 520, 400][m - 1]), (d(15), "Credit card payment", "Credit card", 250),
           (d(15), "Student loan", "Student loan", 150), (d(17), "Cinema", "Entertainment", round(rnd.uniform(20, 45), 2)),
           (d(18), "Weekly shop", "Groceries", round(rnd.uniform(95, 130), 2)), (d(20), "Clothes", "Shopping", round(rnd.uniform(60, 190), 2)),
           (d(22), "Car insurance", "Insurance", 160), (d(24), "Takeaway", "Eating out", round(rnd.uniform(25, 60), 2)),
           (d(25), "Weekly shop", "Groceries", round(rnd.uniform(95, 130), 2)), (d(27), "Fuel", "Transport", round(rnd.uniform(35, 60), 2))]
    if m == 2: ex.append((d(14), "Valentine's gift", "Gifts", 65))
    if m == 3: ex += [(d(9), "Pharmacy", "Health", 28.5), (d(28), "Flights deposit", "Travel", 210)]
for i in range(T0, T1 + 1):
    for col in (2, 3, 4, 5): inp(ws.cell(row=i, column=col), "yyyy-mm-dd" if col == 2 else MONEY if col == 5 else None)
    ws.cell(row=i, column=6, value=f'=IF(D{i}="","",IFERROR(INDEX({GRPS},MATCH(D{i},{CATS},0)),"Check category"))')
    ws.cell(row=i, column=7, value=f'=IF(B{i}="","",MONTH(B{i}))')
    ws.cell(row=i, column=8, value=f'=IF(B{i}="","",YEAR(B{i}))')
    for col in (6, 7, 8):
        c = ws.cell(row=i, column=col); c.font = F(color=MUTED); c.fill = fill("F4F6F5"); c.alignment = Alignment(horizontal="center")
for i, (d, desc, cat, amt) in enumerate(sorted(ex)):
    ws.cell(row=T0 + i, column=2, value=d); ws.cell(row=T0 + i, column=3, value=desc)
    ws.cell(row=T0 + i, column=4, value=cat); ws.cell(row=T0 + i, column=5, value=amt)
ws.cell(row=T0, column=3).comment = Comment("Example rows (January to March) show the format. Delete them when you start your own log.", "Sorted")
dv = DataValidation(type="list", formula1=CATS, allow_blank=True, showErrorMessage=True,
                    errorTitle="Unknown category", error="Pick a category from the list, or add it on the Setup tab first.")
ws.add_data_validation(dv); dv.add(f"D{T0}:D{T1}")
ws.conditional_formatting.add(f"F{T0}:F{T1}", CellIsRule(operator="equal", formula=['"Check category"'], font=Font(name="Arial", color=CORAL, bold=True)))
ws.freeze_panes = "B5"
TD, TA, TG, TM, TY = (f"Transactions!${c}${T0}:${c}${T1}" for c in "BEFGH")
TC = f"Transactions!$D${T0}:$D${T1}"

# ------------------------------------------------------------------ This Month
ws = sheet("This Month", "2E7D5B", {"A": 3, "B": 26, "C": 12, "D": 15, "E": 15, "F": 15, "G": 12, "H": 3, "I": 22, "J": 15, "K": 12, "L": 12})
title(ws, "This Month", "Pick a month. Everything below updates.")
ws["B5"] = "Month"; ws["B5"].font = F(bold=True)
ws["C5"] = "January"; inp(ws["C5"]); ws.merge_cells("C5:D5")
dv = DataValidation(type="list", formula1='"' + ",".join(MONTHS) + '"'); ws.add_data_validation(dv); dv.add("C5")
ws["E5"] = '=MATCH(C5,{"' + '","'.join(MONTHS) + '"},0)'; ws["E5"].font = F(color="FFFFFF")  # month number, hidden in white
ws["F5"] = "=Setup!C5"; ws["F5"].number_format = "0"; ws["F5"].font = F(bold=True)
header(ws, 7, 2, ["Category", "Group", "Budget", "Actual", "Left", "Used"])
M0 = 8
for k in range(30):
    r, s = M0 + k, CAT0 + k
    ws.cell(row=r, column=2, value=f'=IF(Setup!B{s}="","",Setup!B{s})')
    ws.cell(row=r, column=3, value=f'=IF(Setup!B{s}="","",Setup!C{s})')
    ws.cell(row=r, column=4, value=f'=IF(B{r}="","",Setup!D{s})')
    ws.cell(row=r, column=5, value=f'=IF(B{r}="","",SUMIFS({TA},{TC},B{r},{TM},$E$5,{TY},$F$5))')
    ws.cell(row=r, column=6, value=f'=IF(B{r}="","",IF(C{r}="Income",E{r}-D{r},D{r}-E{r}))')
    ws.cell(row=r, column=7, value=f'=IF(OR(B{r}="",N(D{r})=0),"",E{r}/D{r})')
    for col, fm in ((2, None), (3, None), (4, MONEY), (5, MONEY), (6, MONEY), (7, PCT)): out(ws.cell(row=r, column=col), fm)
    ws.cell(row=r, column=3).font = F(color=MUTED)
MR = f"{M0}:{M0+29}"
ws.conditional_formatting.add(f"G{M0}:G{M0+29}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="8CCBA9"))
ws.conditional_formatting.add(f"F{M0}:F{M0+29}", CellIsRule(operator="lessThan", formula=["-0.005"], font=Font(name="Arial", color=CORAL, bold=True)))
# summary
ws["I7"] = "Summary"; ws["I7"].font = F(size=12, bold=True)
summ = [("Income received", f'=SUMIFS({TA},{TG},"Income",{TM},$E$5,{TY},$F$5)'),
        ("Spent (needs + wants)", f'=SUMIFS({TA},{TG},"Needs",{TM},$E$5,{TY},$F$5)+SUMIFS({TA},{TG},"Wants",{TM},$E$5,{TY},$F$5)'),
        ("Saved", f'=SUMIFS({TA},{TG},"Savings",{TM},$E$5,{TY},$F$5)'),
        ("Debt paid", f'=SUMIFS({TA},{TG},"Debt",{TM},$E$5,{TY},$F$5)'),
        ("Left over", "=J8-J9-J10-J11")]
for i, (lab, f) in enumerate(summ):
    ws.cell(row=8 + i, column=9, value=lab).font = F(bold=(i == 4))
    out(ws.cell(row=8 + i, column=10, value=f), MONEY, bold=(i == 4))
ws.conditional_formatting.add("J12", CellIsRule(operator="lessThan", formula=["-0.005"], font=Font(name="Arial", color=CORAL, bold=True)))
ws["I15"] = "50 / 30 / 20 check"; ws["I15"].font = F(size=12, bold=True)
header(ws, 16, 9, ["Group", "% of income", "Target"])
for i, (g, tgt) in enumerate((("Needs", 0.5), ("Wants", 0.3), ("Savings", 0.2))):
    r = 17 + i
    ws.cell(row=r, column=9, value=g).font = F()
    f = f'=IFERROR((SUMIFS({TA},{TG},"{g}",{TM},$E$5,{TY},$F$5)' + (f'+SUMIFS({TA},{TG},"Debt",{TM},$E$5,{TY},$F$5)' if g == "Savings" else "") + ")/$J$8,0)"
    out(ws.cell(row=r, column=10, value=f), PCT)
    out(ws.cell(row=r, column=11, value=tgt), PCT)
ws["I20"] = "Savings includes debt paid down, as in the classic 50/30/20 rule."
ws["I20"].font = F(size=9, italic=True, color=MUTED)
ws.freeze_panes = "B8"

# ------------------------------------------------------------------ Year Dashboard
ws = sheet("Year Dashboard", "1F4D3A", {"A": 3, "B": 14, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13})
title(ws, "Year Dashboard", "Your year in one screen.")
ws["B4"] = "=\"Year \"&Setup!C5"; ws["B4"].font = F(bold=True, color=MUTED)
header(ws, 6, 2, ["Month", "Income", "Needs", "Wants", "Savings", "Debt", "Left over", "Saving rate"])
for i, m in enumerate(MONTHS):
    r = 7 + i
    ws.cell(row=r, column=2, value=m[:3]).font = F(bold=True)
    for j, g in enumerate(GROUPS):
        out(ws.cell(row=r, column=3 + j, value=f'=SUMIFS({TA},{TG},"{g}",{TM},{i+1},{TY},Setup!$C$5)'), MONEY)
    out(ws.cell(row=r, column=8, value=f"=C{r}-SUM(D{r}:G{r})"), MONEY)
    out(ws.cell(row=r, column=9, value=f"=IF(C{r}=0,\"\",(F{r}+G{r})/C{r})"), PCT)
ws.cell(row=19, column=2, value="Total").font = F(bold=True)
for col in "CDEFGH":
    out(ws[f"{col}19"], MONEY, bold=True); ws[f"{col}19"] = f"=SUM({col}7:{col}18)"
ws["I19"] = '=IF(C19=0,"",(F19+G19)/C19)'; out(ws["I19"], PCT, bold=True)
for col in "BCDEFGHI": ws[f"{col}19"].fill = fill(MINT)
ws.conditional_formatting.add("H7:H19", CellIsRule(operator="lessThan", formula=["-0.005"], font=Font(name="Arial", color=CORAL, bold=True)))
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Where the money went, by month"; ch.y_axis.title = None; ch.height, ch.width = 8, 18
ch.add_data(Reference(ws, min_col=4, max_col=7, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=7, max_row=18))
for s, c in zip(ch.series, ["2E7D5B", "E0A33F", "5BA67F", "E0603F"]): s.graphicalProperties.solidFill = c; s.graphicalProperties.line.solidFill = c
ws.add_chart(ch, "B22")
pie = PieChart(); pie.title = "Year so far"; pie.height, pie.width = 8, 10
pie.add_data(Reference(ws, min_col=4, max_col=7, min_row=19, max_row=19), from_rows=True, titles_from_data=False)
pie.set_categories(Reference(ws, min_col=4, max_col=7, min_row=6, max_row=6))
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
ws.add_chart(pie, "H22")
inc = BarChart(); inc.type = "col"; inc.title = "Income vs left over"; inc.height, inc.width = 8, 18
inc.add_data(Reference(ws, min_col=3, min_row=6, max_row=18), titles_from_data=True)
inc.add_data(Reference(ws, min_col=8, min_row=6, max_row=18), titles_from_data=True)
inc.set_categories(Reference(ws, min_col=2, min_row=7, max_row=18))
for s, c in zip(inc.series, ["1F4D3A", "8CCBA9"]): s.graphicalProperties.solidFill = c; s.graphicalProperties.line.solidFill = c
ws.add_chart(inc, "B39")

# ------------------------------------------------------------------ Debt Payoff
ws = sheet("Debt Payoff", "E0603F", {"A": 3, "B": 22, "C": 14, "D": 10, "E": 14, "F": 16, "G": 3, "H": 26, "I": 16})
title(ws, "Debt Payoff", "List debts in the order you want to clear them. Snowball: smallest balance first. Avalanche: highest interest first.")
header(ws, 5, 2, ["Debt", "Balance", "APR", "Min payment", "Paid off in"])
debts = [("Credit card", 3200, 0.229, 95), ("Car loan", 8400, 0.065, 240), ("Student loan", 14500, 0.05, 160)]
D0, ND = 6, 5
for k in range(ND):
    r = D0 + k
    inp(ws.cell(row=r, column=2)); inp(ws.cell(row=r, column=3), MONEY); inp(ws.cell(row=r, column=4), "0.0%"); inp(ws.cell(row=r, column=5), MONEY)
    if k < len(debts):
        for j, v in enumerate(debts[k]): ws.cell(row=r, column=2 + j, value=v)
ws["B12"] = "Extra paid each month"; ws["B12"].font = F(bold=True); ws["C12"] = 150; inp(ws["C12"], MONEY)
ws["B13"] = "First payment month"; ws["B13"].font = F(bold=True); ws["C13"] = dt.date(YEAR, 1, 1); inp(ws["C13"], "mmm yyyy")
ws["B14"] = "Total paid each month"; ws["B14"].font = F(bold=True); ws["C14"] = f"=SUM(E{D0}:E{D0+ND-1})+C12"; out(ws["C14"], MONEY, bold=True)
# schedule: columns owed K..O, pay P..T, balance U..Y, total Z ; rows 40.. (120 months)
S0, NM = 41, 120
OW, PY, BL = ["K", "L", "M", "N", "O"], ["P", "Q", "R", "S", "T"], ["U", "V", "W", "X", "Y"]
ws["B39"] = "Month-by-month schedule (calculated)"; ws["B39"].font = F(size=12, bold=True)
labels = ["Month"] + [f"Owed {i+1}" for i in range(5)] + [f"Pay {i+1}" for i in range(5)] + [f"Balance {i+1}" for i in range(5)] + ["Total balance"]
for j, lab in enumerate(labels):
    c = ws.cell(row=40, column=10 + j, value=lab); c.font = F(bold=True, color="FFFFFF", size=9); c.fill = fill("8A9A94")
ws["J40"].value = "Month"
for col in ["J"] + OW + PY + BL + ["Z"]: ws.column_dimensions[col].width = 11
for t in range(NM):
    r = S0 + t
    ws[f"J{r}"] = f"=EDATE($C$13,{t})"; ws[f"J{r}"].number_format = "mmm yyyy"
    for i in range(5):
        prev = f"N($C${D0+i})" if t == 0 else f"{BL[i]}{r-1}"
        ws[f"{OW[i]}{r}"] = f"={prev}*(1+N($D${D0+i})/12)"
    for i in range(5):
        later = "+".join(f"MIN({OW[j]}{r},N($E${D0+j}))" for j in range(i + 1, 5)) or "0"
        earlier = "+".join(f"{PY[j]}{r}" for j in range(i)) or "0"
        ws[f"{PY[i]}{r}"] = f"=MAX(0,MIN({OW[i]}{r},$C$14-({earlier})-({later})))"
        ws[f"{BL[i]}{r}"] = f"=ROUND({OW[i]}{r}-{PY[i]}{r},2)"
    ws[f"Z{r}"] = f"=SUM(U{r}:Y{r})"
    for col in OW + PY + BL + ["Z"]: ws[f"{col}{r}"].number_format = MONEY; ws[f"{col}{r}"].font = F(size=9, color=MUTED)
    ws[f"J{r}"].font = F(size=9, color=MUTED)
last = S0 + NM - 1
for k in range(ND):
    r = D0 + k
    ws.cell(row=r, column=6, value=f'=IF(N(C{r})=0,"",IF(COUNTIF({BL[k]}{S0}:{BL[k]}{last},">0.005")>={NM},"10+ years",TEXT(EDATE($C$13,COUNTIF({BL[k]}{S0}:{BL[k]}{last},">0.005")),"mmm yyyy")))')
    out(ws.cell(row=r, column=6)); ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
ws["H5"] = "Your plan"; ws["H5"].font = F(size=12, bold=True)
res = [("Debt-free in", f'=IF(COUNTIF(Z{S0}:Z{last},">0.005")>={NM},"10+ years",COUNTIF(Z{S0}:Z{last},">0.005")+1&" months")', None),
       ("Debt-free by", f'=IF(COUNTIF(Z{S0}:Z{last},">0.005")>={NM},"Pay more each month",TEXT(EDATE($C$13,COUNTIF(Z{S0}:Z{last},">0.005")),"mmmm yyyy"))', None),
       ("Starting debt", f"=SUM(C{D0}:C{D0+ND-1})", MONEY),
       ("Total interest you'll pay", f"=SUM(P{S0}:T{last})-I8", MONEY)]
for i, (lab, f, fm) in enumerate(res):
    ws.cell(row=6 + i, column=8, value=lab).font = F(bold=True)
    c = ws.cell(row=6 + i, column=9, value=f); out(c, fm, bold=True); c.alignment = Alignment(horizontal="right")
ws["I6"].font = F(bold=True, size=14, color=BRAND); ws["I7"].font = F(bold=True, size=12, color=BRAND)
ws["H11"] = '=IF(C14<SUM(E6:E10),"Total paid is below your minimum payments. Raise the extra amount.","")'
ws["H11"].font = F(bold=True, color=CORAL)
bal = BarChart(); bal.type = "col"; bal.grouping = "stacked"; bal.overlap = 100; bal.title = "Debt left, month by month"
bal.height, bal.width = 9, 22
bal.add_data(Reference(ws, min_col=21, max_col=23, min_row=40, max_row=S0 + 59), titles_from_data=True)
bal.set_categories(Reference(ws, min_col=10, min_row=S0, max_row=S0 + 59))
for s, c in zip(bal.series, ["E0603F", "E0A33F", "2E7D5B"]): s.graphicalProperties.solidFill = c; s.graphicalProperties.line.solidFill = c
ws.add_chart(bal, "B17")
ws.freeze_panes = "A5"

# ------------------------------------------------------------------ Savings Goals
ws = sheet("Savings Goals", "8CCBA9", {"A": 3, "B": 26, "C": 14, "D": 14, "E": 14, "F": 12, "G": 16, "H": 14})
title(ws, "Savings Goals", "Give every goal a date. Sorted works out the monthly amount.")
header(ws, 5, 2, ["Goal", "Target", "Saved so far", "Target date", "Months left", "Save per month", "Progress"])
# Example target dates are relative to today so they never look stale; customers overwrite them with real dates.
goals = [("Emergency fund (3 months)", 9000, 2400, "=EDATE(TODAY(),12)"), ("Summer holiday", 2500, 600, "=EDATE(TODAY(),7)"),
         ("New laptop", 1400, 350, "=EDATE(TODAY(),5)"), ("House deposit", 30000, 4200, "=EDATE(TODAY(),36)")]
for k in range(12):
    r = 6 + k
    inp(ws.cell(row=r, column=2)); inp(ws.cell(row=r, column=3), MONEY); inp(ws.cell(row=r, column=4), MONEY); inp(ws.cell(row=r, column=5), "mmm yyyy")
    if k < len(goals):
        for j, v in enumerate(goals[k]): ws.cell(row=r, column=2 + j, value=v)
    ws.cell(row=r, column=6, value=f'=IF(OR(B{r}="",E{r}=""),"",MAX(1,(YEAR(E{r})-YEAR(TODAY()))*12+MONTH(E{r})-MONTH(TODAY())))')
    ws.cell(row=r, column=7, value=f'=IF(F{r}="","",MAX(0,(C{r}-D{r})/F{r}))')
    ws.cell(row=r, column=8, value=f'=IF(N(C{r})=0,"",MIN(1,D{r}/C{r}))')
    out(ws.cell(row=r, column=6), "0"); out(ws.cell(row=r, column=7), MONEY, bold=True); out(ws.cell(row=r, column=8), PCT)
ws.conditional_formatting.add("H6:H17", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5BA67F"))
ws["B19"] = "Total to set aside each month"; ws["B19"].font = F(bold=True)
ws["G19"] = "=SUM(G6:G17)"; out(ws["G19"], MONEY, bold=True); ws["G19"].fill = fill(MINT)
ws["B20"] = "Compare this with your Savings budget on the Setup tab."; ws["B20"].font = F(size=9, italic=True, color=MUTED)

for chart_ws in (wb["Year Dashboard"], wb["Debt Payoff"]):
    for c in chart_ws._charts:
        if hasattr(c, "y_axis"): c.y_axis.scaling.min = 0; c.y_axis.numFmt = "#,##0"
for w in wb.worksheets:
    w.sheet_view.zoomScale = 110
    w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
